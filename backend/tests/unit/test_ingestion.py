"""Phase 5 deterministic parsing and structure-aware chunking tests."""

from __future__ import annotations

import io

import pytest
from docx import Document
from openpyxl import Workbook
from pydantic import ValidationError

from app.adapters.ingestion import parse_document
from app.domain.ingestion import (
    CHUNKER_VERSION,
    ChunkLocator,
    DocumentBlock,
    DocumentFormat,
    ParsedDocument,
    StructureKind,
    chunk_document,
)
from app.ports.errors import PermanentPortError
from tests.traceability import req


def _pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body.extend(f"{index} 0 obj\n".encode())
        body.extend(obj)
        body.extend(b"\nendobj\n")
    xref = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    body.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        body.extend(f"{offset:010d} 00000 n \n".encode())
    body.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
        ).encode()
    )
    return bytes(body)


def _docx() -> bytes:
    document = Document()
    document.add_heading("Evidence", level=1)
    document.add_paragraph("Alpha paragraph")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "alpha"
    table.rows[1].cells[1].text = "42"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Evidence"
    sheet.append(["Name", "Value"])
    sheet.append(["alpha", 42])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@req("FR-401")
@pytest.mark.parametrize(
    ("format_", "body", "expected_text"),
    [
        (DocumentFormat.PDF, _pdf("Evidence alpha"), "Evidence alpha"),
        (DocumentFormat.DOCX, _docx(), "Alpha paragraph"),
        (DocumentFormat.TXT, b"Evidence alpha\n\nSecond paragraph", "Evidence alpha"),
        (DocumentFormat.MARKDOWN, b"# Evidence\n\nAlpha paragraph", "Alpha paragraph"),
        (DocumentFormat.CSV, b"name,value\nalpha,42\n", "alpha | 42"),
        (DocumentFormat.XLSX, _xlsx(), "alpha | 42"),
        (DocumentFormat.JSON, b'{"evidence":{"name":"alpha","value":42}}', '"alpha"'),
        (
            DocumentFormat.HTML,
            b"<h1>Evidence</h1><p>Alpha paragraph</p>",
            "Alpha paragraph",
        ),
    ],
    ids=("pdf", "docx", "txt", "markdown", "csv", "xlsx", "json", "html"),
)
def test_all_required_formats_parse_to_stable_chunks(
    format_: DocumentFormat, body: bytes, expected_text: str
) -> None:
    first = parse_document(format_, body)
    second = parse_document(format_, body)

    assert first == second
    assert first.format is format_
    assert first.parser
    assert first.parser_version
    assert any(expected_text in block.text for block in first.blocks)

    chunks = chunk_document(first)
    assert chunks == chunk_document(second)
    assert chunks
    assert [chunk.ordinal for chunk in chunks] == list(range(len(chunks)))
    assert all(chunk.content_hash.startswith("sha256:") for chunk in chunks)
    assert all(chunk.chunker_version == CHUNKER_VERSION for chunk in chunks)


@req("FR-401")
def test_markdown_sections_and_html_table_rows_preserve_structure() -> None:
    markdown = parse_document(DocumentFormat.MARKDOWN, b"# First\n\nOne\n\n## Second\n\nTwo")
    assert [block.locator.section for block in markdown.blocks] == [
        "First",
        "First",
        "Second",
        "Second",
    ]
    assert markdown.blocks[0].kind is StructureKind.HEADING

    html = parse_document(
        DocumentFormat.HTML,
        b"<table><tr><th>Name</th><th>Value</th></tr><tr><td>alpha</td><td>42</td></tr></table>",
    )
    assert [block.kind for block in html.blocks] == [
        StructureKind.TABLE_ROW,
        StructureKind.TABLE_ROW,
    ]
    assert html.blocks[1].repeat_prefix == "Name | Value"
    assert html.blocks[1].locator.table == 1
    assert html.blocks[1].locator.row == 2


@req("FR-401")
def test_chunker_has_exact_overlap_and_byte_stable_locator_hashes() -> None:
    text = " ".join(f"token-{index}" for index in range(20))
    document = ParsedDocument(
        format=DocumentFormat.TXT,
        parser="fixture",
        parser_version="1",
        blocks=(
            DocumentBlock(
                kind=StructureKind.PARAGRAPH,
                text=text,
                locator=ChunkLocator(char_start=10, char_end=10 + len(text), paragraph=1),
            ),
        ),
    )

    chunks = chunk_document(document, target_tokens=8, overlap_tokens=2)

    assert [chunk.text.split() for chunk in chunks] == [
        [f"token-{index}" for index in range(0, 8)],
        [f"token-{index}" for index in range(6, 14)],
        [f"token-{index}" for index in range(12, 20)],
    ]
    assert chunks[1].locator.char_start == 10 + text.index("token-6")
    assert chunks == chunk_document(document, target_tokens=8, overlap_tokens=2)

    moved = document.model_copy(
        update={
            "blocks": (
                document.blocks[0].model_copy(
                    update={
                        "locator": document.blocks[0].locator.model_copy(update={"paragraph": 2})
                    }
                ),
            )
        }
    )
    assert (
        chunks[0].content_hash
        != chunk_document(moved, target_tokens=8, overlap_tokens=2)[0].content_hash
    )


@req("FR-401")
def test_default_policy_targets_800_tokens_with_150_token_overlap() -> None:
    text = " ".join(f"token-{index}" for index in range(1000))
    document = ParsedDocument(
        format=DocumentFormat.TXT,
        parser="fixture",
        parser_version="1",
        blocks=(
            DocumentBlock(
                kind=StructureKind.PARAGRAPH,
                text=text,
                locator=ChunkLocator(char_start=0, char_end=len(text), paragraph=1),
            ),
        ),
    )

    chunks = chunk_document(document)

    assert [chunk.token_count for chunk in chunks] == [800, 350]
    assert chunks[0].text.split()[-150:] == chunks[1].text.split()[:150]


@req("FR-401")
def test_table_rows_stay_whole_and_repeat_header() -> None:
    document = parse_document(
        DocumentFormat.CSV,
        b"name,value\nalpha,the value has more than three tokens\n",
    )

    chunks = chunk_document(document, target_tokens=3, overlap_tokens=1)

    assert len(chunks) == 2
    assert chunks[1].text.startswith("name | value\nalpha | the value")
    assert chunks[1].token_count > 3
    assert chunks[1].locator.row == 2


@req("FR-401")
@pytest.mark.parametrize(
    ("format_", "body"),
    [
        (DocumentFormat.TXT, b""),
        (DocumentFormat.TXT, b"\xff"),
        (DocumentFormat.JSON, b"{"),
        (DocumentFormat.PDF, b"not a pdf"),
        (DocumentFormat.DOCX, b"not a zip"),
        (DocumentFormat.XLSX, b"not a zip"),
        (DocumentFormat.HTML, b"<br>"),
    ],
)
def test_malformed_or_empty_documents_fail_closed(format_: DocumentFormat, body: bytes) -> None:
    with pytest.raises(PermanentPortError, match=format_.value):
        parse_document(format_, body)


@req("FR-401")
@pytest.mark.parametrize(
    ("target", "overlap", "message"),
    [
        (0, 0, "target_tokens must be positive"),
        (10, -1, "overlap_tokens must be non-negative"),
        (10, 10, "overlap_tokens must be non-negative"),
        (10, 11, "overlap_tokens must be non-negative"),
    ],
)
def test_invalid_chunk_policies_are_rejected(target: int, overlap: int, message: str) -> None:
    document = parse_document(DocumentFormat.TXT, b"alpha")
    with pytest.raises(ValueError, match=message):
        chunk_document(document, target_tokens=target, overlap_tokens=overlap)


@req("FR-401")
def test_ingestion_values_are_strict_immutable_and_spans_are_ordered() -> None:
    document = parse_document(DocumentFormat.TXT, b"alpha")
    attribute = "char_start"
    with pytest.raises(ValidationError):
        setattr(document.blocks[0].locator, attribute, 1)
    with pytest.raises(ValidationError, match="char_end"):
        ChunkLocator(char_start=1, char_end=1)
