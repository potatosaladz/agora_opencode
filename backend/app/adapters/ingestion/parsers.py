"""Format-specific parsers normalized into the ingestion domain contract."""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Iterable
from importlib.metadata import version
from typing import Any
from zipfile import BadZipFile

from bs4 import BeautifulSoup
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from app.domain.ingestion import (
    ChunkLocator,
    DocumentBlock,
    DocumentFormat,
    ParsedDocument,
    StructureKind,
)
from app.ports.errors import PermanentPortError

__all__ = ["parse_document"]


# trace: FR-401
def _text(body: bytes) -> str:
    return body.decode("utf-8-sig")


def _block(
    kind: StructureKind,
    text: str,
    *,
    start: int = 0,
    repeat_prefix: str | None = None,
    page: int | None = None,
    section: str | None = None,
    paragraph: int | None = None,
    table: int | None = None,
    row: int | None = None,
    sheet: str | None = None,
    json_path: str | None = None,
) -> DocumentBlock | None:
    line_normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    leading = len(line_normalized) - len(line_normalized.lstrip())
    normalized = line_normalized.strip()
    if not normalized:
        return None
    return DocumentBlock(
        kind=kind,
        text=normalized,
        locator=ChunkLocator(
            page=page,
            section=section,
            paragraph=paragraph,
            table=table,
            row=row,
            sheet=sheet,
            json_path=json_path,
            char_start=start + leading,
            char_end=start + leading + len(normalized),
        ),
        repeat_prefix=repeat_prefix,
    )


def _rows(rows: Iterable[Iterable[Any]], *, sheet: str | None = None) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    header: str | None = None
    for index, row in enumerate(rows, start=1):
        rendered = " | ".join("" if value is None else str(value) for value in row)
        if not rendered.strip(" |"):
            continue
        if header is None:
            header = rendered
        item = _block(
            StructureKind.TABLE_ROW,
            rendered,
            table=1,
            row=index,
            sheet=sheet,
            repeat_prefix=header if index > 1 else None,
        )
        if item is not None:
            blocks.append(item)
    return blocks


def _plain(format_: DocumentFormat, body: bytes) -> ParsedDocument:
    text = _text(body).replace("\r\n", "\n").replace("\r", "\n")
    blocks = []
    offset = 0
    for index, paragraph in enumerate(text.split("\n\n"), start=1):
        item = _block(StructureKind.PARAGRAPH, paragraph, start=offset, paragraph=index)
        if item is not None:
            blocks.append(item)
        offset += len(paragraph) + 2
    return ParsedDocument(
        format=format_, parser="stdlib-text", parser_version="1", blocks=tuple(blocks)
    )


def _markdown(body: bytes) -> ParsedDocument:
    text = _text(body).replace("\r\n", "\n").replace("\r", "\n")
    blocks: list[DocumentBlock] = []
    section: str | None = None
    offset = 0
    for index, paragraph in enumerate(text.split("\n\n"), start=1):
        stripped = paragraph.strip()
        heading = stripped.startswith("#") and stripped.lstrip("#").startswith(" ")
        if heading:
            section = stripped.lstrip("#").strip()
        item = _block(
            StructureKind.HEADING if heading else StructureKind.PARAGRAPH,
            paragraph,
            start=offset,
            paragraph=index,
            section=section,
        )
        if item is not None:
            blocks.append(item)
        offset += len(paragraph) + 2
    return ParsedDocument(
        format=DocumentFormat.MARKDOWN,
        parser="agora-markdown",
        parser_version="1",
        blocks=tuple(blocks),
    )


def _json_blocks(value: Any, path: str = "$") -> list[DocumentBlock]:
    if isinstance(value, dict):
        result: list[DocumentBlock] = []
        for key in sorted(value):
            result.extend(_json_blocks(value[key], f"{path}.{key}"))
        return result
    if isinstance(value, list):
        result = []
        for index, item in enumerate(value):
            result.extend(_json_blocks(item, f"{path}[{index}]"))
        return result
    rendered = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    item = _block(StructureKind.DATA, rendered, json_path=path)
    return [] if item is None else [item]


def _parse_document(format_: DocumentFormat, body: bytes) -> ParsedDocument:
    if format_ in {DocumentFormat.TXT, DocumentFormat.MARKDOWN}:
        return _markdown(body) if format_ is DocumentFormat.MARKDOWN else _plain(format_, body)
    if format_ is DocumentFormat.CSV:
        blocks = _rows(csv.reader(io.StringIO(_text(body), newline="")))
        return ParsedDocument(
            format=format_, parser="stdlib-csv", parser_version="1", blocks=tuple(blocks)
        )
    if format_ is DocumentFormat.JSON:
        value = json.loads(_text(body))
        return ParsedDocument(
            format=format_,
            parser="stdlib-json",
            parser_version="1",
            blocks=tuple(_json_blocks(value)),
        )
    if format_ is DocumentFormat.PDF:
        blocks = []
        for page_number, page in enumerate(PdfReader(io.BytesIO(body)).pages, start=1):
            item = _block(StructureKind.PARAGRAPH, page.extract_text() or "", page=page_number)
            if item is not None:
                blocks.append(item)
        return ParsedDocument(
            format=format_, parser="pypdf", parser_version=version("pypdf"), blocks=tuple(blocks)
        )
    if format_ is DocumentFormat.DOCX:
        document = Document(io.BytesIO(body))
        docx_blocks: list[DocumentBlock] = []
        docx_section: str | None = None
        paragraph_index = 0
        table_index = 0
        for content in document.iter_inner_content():
            if isinstance(content, Paragraph):
                paragraph_index += 1
                is_heading = content.style is not None and content.style.name.startswith("Heading")
                if is_heading:
                    docx_section = content.text.strip()
                item = _block(
                    StructureKind.HEADING if is_heading else StructureKind.PARAGRAPH,
                    content.text,
                    paragraph=paragraph_index,
                    section=docx_section,
                )
                if item is not None:
                    docx_blocks.append(item)
            elif isinstance(content, Table):
                table_index += 1
                rows = _rows([cell.text for cell in row.cells] for row in content.rows)
                docx_blocks.extend(
                    row.model_copy(
                        update={"locator": row.locator.model_copy(update={"table": table_index})}
                    )
                    for row in rows
                )
        return ParsedDocument(
            format=format_,
            parser="python-docx",
            parser_version=version("python-docx"),
            blocks=tuple(docx_blocks),
        )
    if format_ is DocumentFormat.XLSX:
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        xlsx_blocks: list[DocumentBlock] = []
        for sheet in workbook.worksheets:
            xlsx_blocks.extend(_rows(sheet.iter_rows(values_only=True), sheet=sheet.title))
        workbook.close()
        return ParsedDocument(
            format=format_,
            parser="openpyxl",
            parser_version=version("openpyxl"),
            blocks=tuple(xlsx_blocks),
        )
    if format_ is DocumentFormat.HTML:
        soup = BeautifulSoup(_text(body), "html.parser")
        html_blocks: list[DocumentBlock] = []
        html_section: str | None = None
        table_numbers = {
            id(table): index for index, table in enumerate(soup.find_all("table"), start=1)
        }
        row_counts: dict[int, int] = {}
        table_headers: dict[int, str] = {}
        for index, node in enumerate(soup.find_all(["h1", "h2", "h3", "p", "tr"]), start=1):
            if node.name in {"h1", "h2", "h3"}:
                html_section = node.get_text(" ", strip=True)
                kind = StructureKind.HEADING
            elif node.name == "tr":
                kind = StructureKind.TABLE_ROW
            else:
                kind = StructureKind.PARAGRAPH
            if node.name == "tr":
                table_node = node.find_parent("table")
                table_id = id(table_node)
                table_number = table_numbers[table_id]
                row_number = row_counts.get(table_id, 0) + 1
                row_counts[table_id] = row_number
                rendered = " | ".join(
                    cell.get_text(" ", strip=True)
                    for cell in node.find_all(["th", "td"], recursive=False)
                )
                if row_number == 1:
                    table_headers[table_id] = rendered
                item = _block(
                    kind,
                    rendered,
                    table=table_number,
                    row=row_number,
                    section=html_section,
                    repeat_prefix=table_headers.get(table_id) if row_number > 1 else None,
                )
            else:
                item = _block(
                    kind,
                    node.get_text(" ", strip=True),
                    paragraph=index,
                    section=html_section,
                )
            if item is not None:
                html_blocks.append(item)
        return ParsedDocument(
            format=format_,
            parser="beautifulsoup4",
            parser_version=version("beautifulsoup4"),
            blocks=tuple(html_blocks),
        )
    raise ValueError(f"unsupported document format: {format_}")


def parse_document(format_: DocumentFormat, body: bytes) -> ParsedDocument:
    """Parse bytes by declared format and translate malformed input at the adapter boundary."""
    try:
        document = _parse_document(format_, body)
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        PackageNotFoundError,
        PyPdfError,
        UnicodeError,
        ValueError,
    ) as exc:
        raise PermanentPortError(
            f"invalid {format_.value} document", port="document_parser", cause=exc
        ) from exc
    if not document.blocks:
        raise PermanentPortError(
            f"{format_.value} document contains no extractable text", port="document_parser"
        )
    return document
